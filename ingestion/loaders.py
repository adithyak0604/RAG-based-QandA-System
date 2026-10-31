"""
Format-specific loaders. Each one returns a list of "pages" —
plain dicts, not custom classes, so downstream code doesn't need to
know or care which loader produced them.

Every page dict looks like:
{
    "text": str,
    "page_number": int,   # 1-indexed. For DOCX/TXT, "page" = whole doc (1).
    "source_file": str,
}
"""
import os
from pathlib import Path


def load_pdf(filepath: str) -> list[dict]:
    import pdfplumber

    pages = []
    with pdfplumber.open(filepath) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(
                    {
                        "text": text,
                        "page_number": i,
                        "source_file": os.path.basename(filepath),
                    }
                )
    if not pages:
        # No extractable text -> likely a scanned PDF. Flag it rather
        # than silently returning nothing; OCR is a separate add-on.
        raise ValueError(
            f"No extractable text found in {filepath}. "
            "This may be a scanned PDF requiring OCR (not yet implemented)."
        )
    return pages


def load_docx(filepath: str) -> list[dict]:
    import docx

    doc = docx.Document(filepath)
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [
        {
            "text": full_text,
            "page_number": 1,  # DOCX has no reliable page concept pre-render
            "source_file": os.path.basename(filepath),
        }
    ]


def load_pptx(filepath: str) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(filepath)
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slide_text = "\n".join(t for t in texts if t.strip())
        if slide_text.strip():
            pages.append(
                {
                    "text": slide_text,
                    "page_number": i,  # here "page" = slide number
                    "source_file": os.path.basename(filepath),
                }
            )
    return pages


def load_txt(filepath: str) -> list[dict]:
    text = Path(filepath).read_text(encoding="utf-8", errors="ignore")
    return [
        {
            "text": text,
            "page_number": 1,
            "source_file": os.path.basename(filepath),
        }
    ]


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".txt": load_txt,
}


def load_document(filepath: str) -> list[dict]:
    """Dispatch to the correct loader based on file extension."""
    ext = Path(filepath).suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported: {list(LOADERS.keys())}"
        )
    return loader(filepath)
